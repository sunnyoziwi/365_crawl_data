"""Logic dùng chung: đọc file hợp đồng, gọi Claude trích xuất bảng giá, ghi Excel.

Được dùng bởi cả crawl.py (xử lý hàng loạt theo thư mục) và app.py (giao diện Gradio).
"""
import os
import re
from pathlib import Path

import docx
import openpyxl
import pypdf
from anthropic import Anthropic
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

load_dotenv()

ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
MODEL_NAME = "claude-haiku-4-5"
VALID_EXTS = {".pdf", ".docx", ".doc", ".txt"}

_client = None


def get_client() -> Anthropic:
    global _client
    if _client is None:
        if not ANTHROPIC_API_KEY:
            raise RuntimeError("Thiếu ANTHROPIC_API_KEY. Hãy tạo file .env (xem .env.example) và điền API key.")
        _client = Anthropic(api_key=ANTHROPIC_API_KEY.strip())
    return _client


# =============================================================
# SCHEMA TRÍCH XUẤT DỮ LIỆU CHUẨN DẠNG BẢNG
# =============================================================
EXTRACTION_TOOL = {
    "name": "save_contract_rates",
    "description": "Trích xuất bảng giá phòng chi tiết theo từng mùa/khoảng thời gian từ hợp đồng khách sạn.",
    "input_schema": {
        "type": "object",
        "properties": {
            "rates": {
                "type": "array",
                "description": "Danh sách các dòng giá phòng theo từng loại phòng và khoảng ngày.",
                "items": {
                    "type": "object",
                    "properties": {
                        "city": {"type": "string", "description": "Tên thành phố (ví dụ: Hanoi, Hue...)"},
                        "hotel_name": {"type": "string", "description": "Tên khách sạn"},
                        "room_type": {"type": "string", "description": "Hạng phòng"},
                        "capacity": {"type": ["integer", "string"], "description": "Sức chứa tối đa của phòng"},
                        "from_date": {"type": "string", "description": "Ngày bắt đầu (ví dụ: 01-Nov-26)"},
                        "until_date": {"type": "string", "description": "Ngày kết thúc (ví dụ: 03-May-27)"},
                        "single": {"type": ["string", "number"], "description": "Giá phòng Single"},
                        "double": {"type": ["string", "number"], "description": "Giá phòng Double"},
                        "extra_bed": {"type": ["string", "number"], "description": "Giá Extra Bed"},
                        "triple": {"type": ["string", "number"], "description": "Giá phòng Triple"},
                        "quad": {"type": ["string", "number"], "description": "Giá phòng Quad (nếu hợp đồng có)"},
                    },
                    "required": ["hotel_name", "room_type", "from_date", "until_date"],
                },
            }
        },
        "required": ["rates"],
    },
}

# Giữ cố định 11 cột theo chuẩn
COLUMNS = [
    {"key": "city", "header": "City", "align": "left", "is_price": False, "header_color": "000000"},
    {"key": "hotel_name", "header": "Hotel Name", "align": "left", "is_price": False, "header_color": "000000"},
    {"key": "room_type", "header": "Room type", "align": "left", "is_price": False, "header_color": "000000"},
    {"key": "capacity", "header": "Capacity", "align": "center", "is_price": False, "header_color": "486E3C"},
    {"key": "from_date", "header": "From", "align": "center", "is_price": False, "header_color": "000000"},
    {"key": "until_date", "header": "Until", "align": "center", "is_price": False, "header_color": "000000"},
    {"key": "single", "header": "Single", "align": "right", "is_price": True, "header_color": "000000"},
    {"key": "double", "header": "Double", "align": "right", "is_price": True, "header_color": "000000"},
    {"key": "extra_bed", "header": "Extra Bed", "align": "right", "is_price": True, "header_color": "000000"},
    {"key": "triple", "header": "Triple", "align": "right", "is_price": True, "header_color": "000000"},
    {"key": "quad", "header": "Quad", "align": "right", "is_price": True, "header_color": "000000"},
]


def clean_cell_value(val, is_price=False):
    """Nếu không có dữ liệu, hoặc là NA/unknown thì để trống hoàn toàn."""
    if val is None:
        return ""
    val_str = str(val).strip()

    if val_str.lower() in ["", "none", "unknown", "null", "undefined", "na", "n/a", "-"]:
        return ""

    if is_price:
        cleaned = re.sub(r"[^\d.]", "", val_str)
        try:
            if "." in cleaned:
                return float(cleaned)
            return int(cleaned)
        except ValueError:
            return ""

    return val_str


def write_excel(all_rows: list, output_path: Path):
    if not all_rows:
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"

    ws.append([col["header"] for col in COLUMNS])

    for item in all_rows:
        cleaned = {col["key"]: clean_cell_value(item.get(col["key"]), col["is_price"]) for col in COLUMNS}

        # Nếu hợp đồng không có sẵn giá Triple, suy ra = Double + Extra Bed (khi có đủ 2 giá này)
        if cleaned["triple"] == "" and isinstance(cleaned["double"], (int, float)) and isinstance(cleaned["extra_bed"], (int, float)):
            cleaned["triple"] = cleaned["double"] + cleaned["extra_bed"]

        row_data = [cleaned[col["key"]] for col in COLUMNS]
        ws.append(row_data)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11, color="000000")

    cell_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    ws.row_dimensions[1].height = 24
    for col_idx, col_cfg in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = PatternFill(start_color=col_cfg["header_color"], end_color=col_cfg["header_color"], fill_type="solid")
        cell.alignment = Alignment(horizontal=col_cfg["align"], vertical="center")
        cell.border = cell_border

    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 20
        for c, col_cfg in enumerate(COLUMNS, 1):
            cell = ws.cell(row=r, column=c)
            cell.font = data_font
            cell.border = cell_border
            cell.alignment = Alignment(horizontal=col_cfg["align"], vertical="center")

            if col_cfg["is_price"] and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

    ws.auto_filter.ref = ws.dimensions

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if cell.number_format == "#,##0" and isinstance(cell.value, (int, float)):
                val = f"{cell.value:,.0f}"
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    wb.save(output_path)


def read_text_from_file(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif ext == ".docx":
        try:
            doc = docx.Document(file_path)
            full_text = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    full_text.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(full_text)
        except Exception as e:
            print(f"      [!] Lỗi mở docx: {e}")
            return ""
    elif ext == ".pdf":
        try:
            reader = pypdf.PdfReader(str(file_path))
            return "\n".join([page.extract_text() or "" for page in reader.pages])
        except Exception as e:
            print(f"      [!] Lỗi mở pdf: {e}")
            return ""
    elif ext == ".doc":
        try:
            with open(file_path, "rb") as f:
                content = f.read()
            return "".join([chr(b) for b in content if 32 <= b <= 126 or b in (10, 13)])
        except Exception:
            return ""
    return ""


def extract_rates(doc_text: str, hotel_name_hint: str = "") -> list:
    """Gọi Claude để trích xuất danh sách dòng giá từ nội dung văn bản của 1 tài liệu."""
    prompt = (
        f"Bạn là chuyên gia trích xuất dữ liệu bảng giá khách sạn.\n"
        f"Tên khách sạn gợi ý: '{hotel_name_hint}'.\n"
        "Nhiệm vụ: Trích xuất toàn bộ các dòng giá phòng theo từng mùa/khoảng thời gian vào tool save_contract_rates.\n"
        "Mỗi khoảng thời gian của một loại phòng tạo thành một dòng riêng biệt gồm: "
        "City, Hotel Name, Room type, Capacity, From, Until, Single, Double, Extra Bed, Triple, Quad.\n"
        "Quy tắc dữ liệu:\n"
        "- Nếu bất kỳ trường nào không có thông tin (như không có giá Extra Bed, Triple, Quad, Single...), "
        "hãy ĐỂ TRỐNG hoặc bỏ qua trường đó. TUYỆT ĐỐI KHÔNG điền 'NA', 'N/A', 'none', hay 'unknown'.\n"
        "- Quad là giá phòng cho 4 người, chỉ một số hợp đồng có mục này — nếu không thấy trong tài liệu thì để trống, "
        "KHÔNG được tự suy ra hay tính toán giá Quad.\n"
        "- Hợp đồng thường có 2 loại giá: FIT (khách lẻ/Free Independent Traveler) và GIT (khách đoàn/Group Inclusive Tour). "
        "CHỈ trích xuất giá FIT, TUYỆT ĐỐI KHÔNG lấy giá GIT/Group. Nếu tài liệu không ghi rõ FIT/GIT thì coi như bảng giá đó là FIT.\n\n"
        f"--- BẮT ĐẦU TÀI LIỆU ---\n{doc_text}\n--- KẾT THÚC TÀI LIỆU ---"
    )

    response = get_client().messages.create(
        model=MODEL_NAME,
        max_tokens=4096,
        tools=[EXTRACTION_TOOL],
        tool_choice={"type": "tool", "name": "save_contract_rates"},
        messages=[{"role": "user", "content": prompt}],
    )

    for block in response.content:
        if block.type == "tool_use" and block.name == "save_contract_rates":
            return block.input.get("rates", [])
    return []
